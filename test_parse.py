"""第一个测试文件：拿真实数据 + 一条人造边界数据，验证 parse.py 的规矩。

跑：python -m pytest test_parse.py -v
"""

import pytest

import main
import parse

SAMPLE = open("samples/df-h-ldplayer14.txt", encoding="utf-8").read()


def test_row_count():
    """真实输出里有 14 条数据行（表头不算）。"""
    assert len(parse.parse_df(SAMPLE)) == 14


def test_each_row_has_six_fields():
    """每条记录都必须是 6 个字段，一个例外都不能有。"""
    assert {len(row) for row in parse.parse_df(SAMPLE)} == {6}


def test_whitelist_keeps_only_data():
    """白名单只留设备自己的存储：跑完筛完，只剩 /data 这一行。"""
    kept, _ = parse.pick(parse.parse_df(SAMPLE))
    assert [row[-1] for row in kept] == ["/data"]


def test_dropped_rows_have_reason():
    """被排除的行必须带原因，不能静默消失。"""
    _, dropped = parse.pick(parse.parse_df(SAMPLE))
    assert len(dropped) == 13
    assert all(reason for _, reason in dropped)


def test_mount_point_with_space():
    """人造边界：挂载点里带空格时，最后一列不能被切碎。

    模拟器的挂载点都没有空格，所以这条是唯一能测出 maxsplit 是否写对的用例。
    """
    fake = "Filesystem Size Used Avail Use% Mounted on\n"
    fake += "/dev/block/sdb2  23G  154M  22G  1%  /mnt/media_rw/My Disk"
    rows = parse.parse_df(fake)
    assert len(rows) == 1
    assert rows[0][-1] == "/mnt/media_rw/My Disk"


def test_humanize_basic():
    """常规值：单位换对了就行。"""
    assert parse.humanize(512) == "512K"
    assert parse.humanize(2048) == "2.0M"
    assert parse.humanize(1048576) == "1.0G"


def test_humanize_at_boundary():
    """边界值：1048570K = 1023.99M，四舍五入会显示成 1024.0M —— 必须进位成 G。

    这个用例是回归测试：它锁定"判断的对象必须等于显示的对象"这条规则。
    """
    assert parse.humanize(1048570) == "1.0G"
    # 差一点点、但四舍五入后仍然够进位的那个点
    assert parse.humanize(1048525) == "1.0G"
    # 再差一点就不该进位了（1023.9M）
    assert parse.humanize(1048524) == "1023.9M"


# ---------------------------------------------------------------
# 下面两条是"规格"，等你把 parse.match() 写出来让它们变绿
# ---------------------------------------------------------------

def test_match_counts_hits_and_misses():
    """清单 3 条、数据里只存在 2 条 → 命中 2，未命中 1（且名字对得上）。"""
    rows = parse.parse_df(SAMPLE)
    hits, misses = parse.match(rows, ["/data", "/dev", "/nope"])
    assert len(hits) == 2
    assert misses == ["/nope"]


def test_match_keeps_whole_row():
    """命中时要把整行原样带出来，不是只给路径。"""
    rows = parse.parse_df(SAMPLE)
    hits, misses = parse.match(rows, ["/data"])
    assert misses == []
    path, row = hits[0]
    assert path == "/data"
    assert len(row) == 6

# ---------------------------------------------------------------
# 第四格的规格：parse_du()（du 的输出 = 数值 + 路径，两段）
# ---------------------------------------------------------------

DU_SAMPLE = open("samples/du-data.txt", encoding="utf-8").read()


def test_parse_du_two_fields():
    """du 的每一行都必须拆成 2 段：数值 + 路径。"""
    rows = parse.parse_du(DU_SAMPLE)
    assert len(rows) > 1000
    assert {len(r) for r in rows} == {2}


def test_parse_du_value_is_digits():
    """第一段是纯数字（单位 K），不允许出现单位字符。"""
    rows = parse.parse_du(DU_SAMPLE)
    assert all(r[0].isdigit() for r in rows)


def test_parse_du_keeps_path_intact():
    """第二段是完整路径，保留前导 /（归一化是下一步的事，不是解析的事）。"""
    rows = parse.parse_du("37908\t/data/local/file-cache")
    assert rows[0] == ["37908", "/data/local/file-cache"]


# ---------------------------------------------------------------
# 第五格的规格：归一化（清单不带前导 /，du 输出带前导 /）
# ---------------------------------------------------------------

def test_match_normalizes_leading_slash():
    """清单写 `data/fonts`，数据里是 `/data/fonts` —— 也必须匹配上。

    注意：返回的 path 应该保持清单里的原样（data/fonts），不要偷偷改数据。
    """
    rows = parse.parse_du("100\t/data/fonts")
    hits, misses = parse.match(rows, ["data/fonts"])
    assert misses == []
    assert hits[0][0] == "data/fonts"


def test_match_still_exact_on_other_parts():
    """归一化只剥前导 /，不能把 `data/fonts` 和 `data/fonts/files` 混为一谈。"""
    rows = parse.parse_du("100\t/data/fonts/files")
    hits, misses = parse.match(rows, ["data/fonts"])
    assert hits == []
    assert misses == ["data/fonts"]

def test_key_strips_only_leading_slash():
    assert parse._key("/data/fonts") == "data/fonts"
    assert parse._key("data/fonts") == "data/fonts"
    assert parse._key("/data/log/") == "data/log/"     # 尾斜杠不动,归一化


# ---------------------------------------------------------------
# 第六格的规格：load_paths()（读清单：跳过 # 注释和空行）
# ---------------------------------------------------------------

def test_load_paths_skips_comments_and_blanks():
    paths = parse.load_paths("paths.example.txt")
    assert paths, "清单不能为空"
    assert all(not p.startswith("#") for p in paths)
    assert all(p.strip() for p in paths), "空行必须被跳过"
    assert paths[0] == "data/fonts"      # 第一个非注释、非空行
    assert "data" in paths

# ---------------------------------------------------------------
# 练习：top_n()（排序 + lambda）
# ---------------------------------------------------------------

def test_top_n_orders_by_size_desc():
    """按大小从大到小；返回 (数字, 路径) 的元组。"""
    rows = parse.parse_du("100\t/data/a\n900\t/data/b\n50\t/data/c")
    assert parse.top_n(rows, 2) == [(900, "/data/b"), (100, "/data/a")]


def test_top_n_sorts_as_numbers_not_strings():
    """关键的坑：字符串排序时 "9" > "100"，数字排序时 100 > 9。"""
    rows = parse.parse_du("9\t/data/small\n100\t/data/big")
    assert parse.top_n(rows, 1) == [(100, "/data/big")]

# ---------------------------------------------------------------
# 第七格：growth()（两次采集 → 增长%）
# ---------------------------------------------------------------

def _rows(*pairs):
    """小工具：_rows(("100","/data/a"), ...) → [["100","/data/a"], ...]"""
    return [[str(k), p] for k, p in pairs]


def test_growth_basic():
    """正常增长：100K → 150K = +50%。"""
    before = _rows((100, "/data/a"))
    after = _rows((150, "/data/a"))
    assert parse.growth(before, after) == [("/data/a", 100, 150, 50.0)]


def test_growth_decrease():
    """下降要能给负数，不能卡在 0。"""
    before = _rows((200, "/data/a"))
    after = _rows((150, "/data/a"))
    assert parse.growth(before, after) == [("/data/a", 200, 150, -25.0)]


def test_growth_before_zero_is_none():
    """关键：before = 0 时 pct 必须是 None（不可计算），不能崩、也不能编一个数。"""
    before = _rows((0, "/data/a"))
    after = _rows((100, "/data/a"))
    result = parse.growth(before, after)
    assert result == [("/data/a", 0, 100, None)]


def test_growth_missing_on_either_side():
    """只在一边出现的路径也要占一行，另一侧写 None —— 不许静默丢掉。"""
    before = _rows((100, "/data/a"))
    after = _rows((50, "/data/b"))
    result = parse.growth(before, after)
    assert result == [
        ("/data/a", 100, None, None),   # 跑机后消失了
        ("/data/b", None, 50, None),    # 跑机后新出现
    ]

# ---------------------------------------------------------------
# Excel 输出：把增长表落成 xlsx
# ---------------------------------------------------------------

def test_write_excel(tmp_path):
    """写出后重新读回来验证：表头 + 每行 + None 写空单元格。"""
    pytest.importorskip("openpyxl")
    out = tmp_path / "t.xlsx"
    main.write_excel([("data", 100, 150, 50.0), ("data/x", None, 5, None)], str(out))

    from openpyxl import load_workbook

    ws = load_workbook(out).active
    assert ws.max_row == 3                 # 表头 + 2 行
    assert ws.cell(2, 1).value == "data"
    assert ws.cell(2, 4).value == 50.0
    assert ws.cell(3, 2).value is None     # 前值缺失 → 空单元格
    assert ws.cell(3, 4).value is None     # 增长不可计算 → 空单元格

def test_latest_before_picks_newest(tmp_path, monkeypatch):
    """两次 before 快照，取修改时间最新的那个——跨天也能配上对。"""
    import json
    import os

    monkeypatch.setattr(main, "OUT_DIR", str(tmp_path))
    monkeypatch.setattr(main, "SERIAL", "dev")
    folder = tmp_path / "dev"
    folder.mkdir()

    for name, when in (("2026-10-20-before.json", 1000), ("2026-10-21-before.json", 2000)):
        p = folder / name
        p.write_text(json.dumps({"stage": "before", "saved_at": name, "rows": []}), encoding="utf-8")
        os.utime(p, (when, when))

    got = main.latest_before()
    assert got is not None
    assert "2026-10-21" in got["saved_at"]

# ---------------------------------------------------------------
# 使用率：不许用 Use% 列，必须自己算（来自真实 df -k 数据）
# ---------------------------------------------------------------

DF_K_ROW = ["/dev/block/sdb2", "24170284", "182100", "23988184", "1%",
            "/system/vendor/lib/arm/cpuinfo"]


def test_usage_ignores_rounded_use_percent():
    """真实数据：df 显示 1%，真值 0.75% —— 必须按 used/total 算，不能取 row[4]。"""
    total, used, pct = parse.usage(DF_K_ROW)
    assert total == 24170284
    assert used == 182100
    assert pct == 0.75          # 不是 1.0，也不是 0.8


def test_judge_under_rated():
    assert "未超" in parse.judge(50.0, 80.0)


def test_judge_over_rated():
    assert "超" in parse.judge(90.0, 80.0)


# 边界那一条你自己写：pct == rated 时期望什么？
def test_judge_at_boundary():
    """等于额定值不算超标（我们的定义）—— 所以提示语里必须是 ≤。"""
    assert parse.judge(80.0, 80.0).startswith("未超")


def test_judge_keeps_rated_precision():
    """额定值 0.5 不能显示成 0%——判断用的是 0.5，显示也必须是 0.5。"""
    assert "0.5%" in parse.judge(0.75, 0.5)

# ---------------------------------------------------------------
# 内存：free -m 的 Mem 行（Android 是 4 行，不是 3 行）
# ---------------------------------------------------------------

FREE_SAMPLE = """              total        used        free      shared     buffers
Mem:           5944        1037        4906           4           2
-/+ buffers/cache:        1035        4908
Swap:             0           0           0"""


def test_parse_free_reads_mem_line():
    """取 Mem: 行的 total 和 used（M）——不是第 2 行，也不是最后一行。"""
    total, used = parse.parse_free(FREE_SAMPLE)
    assert total == 5944
    assert used == 1037


def test_parse_free_raises_when_missing():
    """找不到 Mem: 行时必须报错，不能返回 (0, 0) 冒充。"""
    with pytest.raises(ValueError):
        parse.parse_free("Swap:             0           0           0")
